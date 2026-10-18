#!/usr/bin/env python3
"""Fetch karuta tournament notices and verify deadlines against linked primary documents."""
import io
import json
import re
from datetime import date, datetime, timedelta
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

LIST_URL = "https://www.karuta.or.jp/cup-info/"
OUT = "data/tournaments.json"
HEADERS = {"User-Agent": "Kokudai-Karuta-Tournament-Bot/1.1"}
MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MAX_DOCUMENTS_PER_EVENT = 8
SESSION = requests.Session()
SESSION.headers.update(HEADERS)


def clean(s):
    return re.sub(r"[ \t\r\n\u3000]+", " ", s or "").strip()


def iso_date(s):
    if not s:
        return None
    m = re.search(r"(20[0-9]{2})[./年-]\s*([0-9]{1,2})[./月-]\s*([0-9]{1,2})日?", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return None
    m = re.search(r"令和\s*([元0-9]+)\s*年\s*([0-9]{1,2})\s*月\s*([0-9]{1,2})\s*日?", s)
    if m:
        year = 2018 + (1 if m.group(1) == "元" else int(m.group(1)))
        try:
            return date(year, int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return None
    return None


def date_candidates(text, event_date):
    """Parse western-calendar, Reiwa-era, and month/day dates near deadline wording."""
    found = []
    pattern = r"令和\s*[元0-9]+\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日?|20[0-9]{2}\s*年\s*[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日?|[0-9]{1,2}\s*月\s*[0-9]{1,2}\s*日?"
    for m in re.finditer(pattern, text):
        token = m.group(0)
        parsed = iso_date(token)
        if not parsed:
            md = re.search(r"([0-9]{1,2})\s*月\s*([0-9]{1,2})", token)
            if not md:
                continue
            year = date.fromisoformat(event_date).year
            try:
                candidate = date(year, int(md.group(1)), int(md.group(2)))
            except ValueError:
                continue
            event = date.fromisoformat(event_date)
            if candidate > event:
                try:
                    candidate = date(year - 1, candidate.month, candidate.day)
                except ValueError:
                    continue
            parsed = candidate.isoformat()
        if parsed not in found:
            found.append(parsed)
    return found


def after_label(lines, label):
    for i, x in enumerate(lines):
        if label not in x:
            continue
        same = clean(x.split(label, 1)[1])
        if same:
            return same
        for y in lines[i + 1:i + 6]:
            y = clean(y)
            if y and label not in y:
                return y
    return ""


def prefecture(text):
    prefs = ["北海道","青森県","岩手県","宮城県","秋田県","山形県","福島県","茨城県","栃木県","群馬県","埼玉県","千葉県","東京都","神奈川県","新潟県","富山県","石川県","福井県","山梨県","長野県","岐阜県","静岡県","愛知県","三重県","滋賀県","京都府","大阪府","兵庫県","奈良県","和歌山県","鳥取県","島根県","岡山県","広島県","山口県","徳島県","香川県","愛媛県","高知県","福岡県","佐賀県","長崎県","熊本県","大分県","宮崎県","鹿児島県","沖縄県"]
    for p in prefs:
        if p in text:
            return p
    city_map = {"横浜市":"神奈川県","川崎市":"神奈川県","相模原市":"神奈川県","鎌倉市":"神奈川県","藤沢市":"神奈川県","平塚市":"神奈川県","さいたま市":"埼玉県","川口市":"埼玉県","川口":"埼玉県","春日部市":"埼玉県","千葉市":"千葉県","成田市":"千葉県","成田":"千葉県","船橋市":"千葉県","柏市":"千葉県","松戸市":"千葉県","市川市":"千葉県","墨田区":"東京都","渋谷区":"東京都","世田谷区":"東京都","新宿区":"東京都","豊島区":"東京都","江東区":"東京都","江戸川区":"東京都","八王子市":"東京都","立川市":"東京都","町田市":"東京都","関東第一":"東京都","関東第一高等学校":"東京都","宇都宮市":"栃木県","宇都宮":"栃木県","水戸市":"茨城県","高崎市":"群馬県","前橋市":"群馬県","大宮":"埼玉県"}
    for c, p in city_map.items():
        if c in text:
            return p
    return ""


def extract_document_text(url):
    """Extract text from supported primary documents; return empty text on unsupported/bad files."""
    try:
        response = SESSION.get(url, timeout=25)
        response.raise_for_status()
        if len(response.content) > MAX_DOCUMENT_BYTES:
            print("document too large; skip", url)
            return ""
        path = url.split("?", 1)[0].lower()
        data = io.BytesIO(response.content)
        if path.endswith(".pdf"):
            reader = PdfReader(data)
            return "\n".join(page.extract_text() or "" for page in reader.pages[:80])
        if path.endswith(".docx"):
            doc = Document(data)
            chunks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    chunks.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(chunks)
        if path.endswith(".xlsx") or path.endswith(".xlsm"):
            book = load_workbook(data, read_only=True, data_only=True)
            chunks = []
            for sheet in book.worksheets[:8]:
                for row in sheet.iter_rows(max_row=250, max_col=20, values_only=True):
                    value = " ".join(str(v) for v in row if v is not None)
                    if value:
                        chunks.append(value)
            book.close()
            return "\n".join(chunks)
    except Exception as exc:
        print("document text extraction skipped", url, str(exc)[:180])
    return ""


def find_deadline(text, event_date):
    """Accept only dates explicitly tied to advance application, never event-day reception."""
    text = clean(text).replace("：", ":")
    patterns = [
        r"(?:参加)?(?:申込|申し込み|申込み)(?:の)?(?:締切|締め切り|期限|締切日|期限日)",
        r"(?:申込|申し込み|申込み)(?:受付)?(?:期間|期限|締切|締め切り)",
        r"(?:申込書|参加申込書)[^。\n]{0,35}(?:必着|締切|締め切り|期限)",
        r"(?:申込|申し込み|申込み)[^。\n]{0,35}(?:まで|必着|消印有効)",
        r"申込期間",
    ]
    hits = []
    for pattern in patterns:
        for match in re.finditer(pattern, text):
            start = max(0, match.start() - 25)
            end = min(len(text), match.end() + 100)
            context = text[start:end]
            after = date_candidates(text[match.end():min(len(text), match.end() + 100)], event_date)
            candidates = after or date_candidates(context, event_date)
            for chosen in candidates:
                parsed = date.fromisoformat(chosen)
                event = date.fromisoformat(event_date)
                # Application must close before the event; ignore reception/check-in on event day.
                if parsed < event and (event - parsed).days <= 370:
                    hits.append(((event - parsed).days, chosen, context))
                    break
    if not hits:
        return None, ""
    hits.sort(key=lambda item: item[0])
    return hits[0][1], hits[0][2]


def parse_detail(url):
    response = SESSION.get(url, timeout=20)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    lines = [clean(x) for x in soup.stripped_strings if clean(x)]
    name = after_label(lines, "大会名称") or (soup.find("h1").get_text(" ", strip=True) if soup.find("h1") else "")
    event_date = iso_date(after_label(lines, "開催日"))
    venue = after_label(lines, "会場")
    rank_text = after_label(lines, "級")
    ranks = sorted(set(re.findall(r"(?<![A-Z])[ABCDE](?=級|[,、 ]|$)", rank_text)), key="ABCDE".index)

    # First try the official event page's explicit application field.
    deadline_text = after_label(lines, "事前申込")
    deadline = iso_date(deadline_text)
    deadline_source_url = url if deadline else ""
    deadline_source_type = "official_page" if deadline else "not_found"

    docs = []
    seen_docs = set()
    file_exts = (".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm")
    keywords = ("大会情報", "大会案内", "開催案内", "参加申込", "申込書", "案内")
    for a in soup.find_all("a", href=True):
        label = clean(a.get_text(" ", strip=True))
        href = urljoin(url, a["href"])
        href_lower = href.lower().split("?", 1)[0]
        if (href_lower.endswith(file_exts) or any(k in label for k in keywords)) and href not in seen_docs:
            seen_docs.add(href)
            if label == "大会情報":
                continue
            docs.append({"label": label or "大会資料", "url": href})

    # If the page has no deadline, inspect linked primary materials (guide/application form).
    if not deadline and event_date:
        preferred = [d for d in docs if any(k in d["label"] for k in ("大会情報案内", "大会案内", "開催案内", "参加申込書", "申込書", "案内"))]
        for doc in preferred[:MAX_DOCUMENTS_PER_EVENT]:
            text = extract_document_text(doc["url"])
            if not text:
                continue
            candidate, evidence = find_deadline(text, event_date)
            if candidate:
                deadline = candidate
                deadline_source_url = doc["url"]
                deadline_source_type = "primary_document"
                print("deadline found in primary document:", name, candidate, doc["url"], evidence[:100])
                break

    series = re.sub(r"第 *[0-9]+ *回", "", name).strip()
    series = re.sub(r"第 *[0-9]+ *次", "", series).strip()
    return {
        "id": "karuta-" + re.sub(r"[^0-9a-z]+", "-", url.lower()).strip("-"),
        "seriesId": series or name,
        "name": name,
        "date": event_date,
        "venue": venue,
        "prefecture": prefecture(" ".join(lines)),
        "ranks": ranks,
        "officialDeadline": deadline,
        "kokudaiDeadline": (date.fromisoformat(deadline) - timedelta(days=7)).isoformat() if deadline else None,
        "deadlineSourceUrl": deadline_source_url or None,
        "deadlineSourceType": deadline_source_type,
        "documents": docs,
        "sourceUrl": url,
        "fetchedAt": datetime.now().astimezone().isoformat(),
    }


def main():
    response = SESSION.get(LIST_URL, timeout=20)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    urls = []
    for a in soup.find_all("a", href=True):
        url = urljoin(LIST_URL, a["href"])
        if "/cup-info/" in url and url.rstrip("/") != LIST_URL.rstrip("/") and "/date/" not in url and url not in urls:
            urls.append(url)

    rows = []
    for url in urls:
        try:
            item = parse_detail(url)
            if item["date"] and item["name"]:
                rows.append(item)
        except Exception as exc:
            print("skip", url, str(exc)[:180])

    today = date.today()
    limit = today + timedelta(days=365)
    rows = [x for x in rows if today.isoformat() <= x["date"] <= limit.isoformat()]
    rows.sort(key=lambda x: x["date"])
    if not rows:
        raise RuntimeError("No tournaments were fetched; refusing to overwrite tournaments.json")

    with open(OUT, "w", encoding="utf-8", newline="") as output:
        json.dump(rows, output, ensure_ascii=False, indent=2)
        output.write("\n")
    with open(OUT, encoding="utf-8") as source:
        loaded = json.load(source)
    if not isinstance(loaded, list) or not all(isinstance(x, dict) and x.get("name") and x.get("date") for x in loaded):
        raise RuntimeError("Generated tournaments.json failed validation")
    verified = sum(1 for x in loaded if x.get("deadlineSourceType") == "primary_document")
    missing = sum(1 for x in loaded if not x.get("officialDeadline"))
    print(f"saved {len(loaded)} tournaments; {verified} deadlines verified from primary documents; {missing} deadlines still need manual confirmation")


if __name__ == "__main__":
    main()
